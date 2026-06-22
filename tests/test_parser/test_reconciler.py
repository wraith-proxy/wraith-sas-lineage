"""Tests for the Excel reconciler."""

from pathlib import Path

from sas_lineage.models import ReconciliationResult
from sas_lineage.parser.reconciler import reconcile

_COLS = {"col_src": "col_mart", "col_src2": "col_mart2"}


def test_no_excel_returns_not_supplied():
    result = reconcile(_COLS, None)
    assert isinstance(result, ReconciliationResult)
    assert result.excel_supplied is False


def test_missing_excel_file(tmp_path: Path):
    result = reconcile(_COLS, tmp_path / "nonexistent.xlsx")
    assert result.excel_supplied is False


def test_excel_supplied_creates_result(tmp_path: Path):
    try:
        import openpyxl
    except ImportError:
        return  # skip if not installed

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["source", "target", "job"])
    ws.append(["col_src", "col_mart", "job1"])
    path = tmp_path / "mapping.xlsx"
    wb.save(path)

    result = reconcile(_COLS, path)
    assert result.excel_supplied is True


def test_excel_match_counted(tmp_path: Path):
    try:
        import openpyxl
    except ImportError:
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["source", "target", "job"])
    ws.append(["col_src", "col_mart", "j1"])
    path = tmp_path / "mapping.xlsx"
    wb.save(path)

    result = reconcile(_COLS, path)
    assert result.matches >= 1


def test_excel_only_detected(tmp_path: Path):
    try:
        import openpyxl
    except ImportError:
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["source", "target", "job"])
    ws.append(["col_not_in_parser", "col_mart_x", "j1"])
    path = tmp_path / "mapping.xlsx"
    wb.save(path)

    result = reconcile(_COLS, path)
    assert any("col_not_in_parser" in e for e in result.excel_only)


def test_mismatch_same_source_different_target(tmp_path: Path):
    try:
        import openpyxl
    except ImportError:
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["source", "target", "job"])
    # Same source as parser, but Excel maps it to a different target.
    ws.append(["col_src", "col_mart_OTHER", "j1"])
    path = tmp_path / "mapping.xlsx"
    wb.save(path)

    result = reconcile(_COLS, path)
    assert any("col_src" in m for m in result.mismatches)
    assert not any("col_src->col_mart " in p for p in result.parser_only)


def test_parser_only_when_no_excel_overlap(tmp_path: Path):
    try:
        import openpyxl
    except ImportError:
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["source", "target", "job"])
    path = tmp_path / "empty.xlsx"
    wb.save(path)

    result = reconcile(_COLS, path)
    assert len(result.parser_only) == len(_COLS)
